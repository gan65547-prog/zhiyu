"""
文档处理器

提供文档文本提取和分块功能
"""

import re
import shutil
from typing import List, Optional, Dict, Any
from abc import ABC, abstractmethod

from core.rag_config import get_rag_config
from infrastructure.logging.logger import get_logger

logger = get_logger(__name__)


class TextSplitter(ABC):
    """文本分块器抽象接口"""

    @abstractmethod
    def split(self, text: str, chunk_size: int = 500, chunk_overlap: int = 50) -> List[str]:
        """分割文本"""
        pass


class RecursiveCharacterTextSplitter(TextSplitter):
    """递归字符文本分块器"""

    def __init__(self, separators: Optional[List[str]] = None):
        self.separators = separators or ["\n\n", "\n", "。", "！", "？", "，", " ", ""]

    def split(self, text: str, chunk_size: int = 1000, chunk_overlap: int = 150) -> List[str]:
        """分割文本"""
        chunks = []
        start = 0

        while start < len(text):
            end = start + chunk_size

            for separator in self.separators:
                if separator and start + chunk_size < len(text):
                    pos = text.rfind(separator, start + chunk_size - 20, end)
                    if pos != -1:
                        end = pos + len(separator)
                        break

            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)

            start = end - chunk_overlap
            if start < 0:
                start = 0

        return chunks


class DocumentProcessor:
    """文档处理器"""

    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 150):
        self.config = get_rag_config()
        self.chunk_size = chunk_size or self.config.vector_size
        self.chunk_overlap = chunk_overlap
        self.text_splitter = RecursiveCharacterTextSplitter()

    @staticmethod
    def ocr_runtime_status() -> Dict[str, Any]:
        """Return a truthful local readiness check for image OCR."""
        missing: List[str] = []
        try:
            import PIL  # noqa: F401
        except ImportError:
            missing.append("Pillow")
        try:
            import pytesseract
        except ImportError:
            missing.append("pytesseract")
            return {"ready": False, "missing": missing}

        executable = shutil.which("tesseract")
        if not executable:
            missing.append("tesseract-ocr")
            return {"ready": False, "missing": missing}
        try:
            languages = set(pytesseract.get_languages(config=""))
        except Exception as exc:  # pragma: no cover - depends on native binary
            logger.warning("OCR 运行时探测失败 executable=%s error=%s", executable, exc)
            return {"ready": False, "missing": [*missing, "tesseract-language-data"]}
        if "chi_sim" not in languages:
            missing.append("tesseract-ocr-chi-sim")
        if "eng" not in languages:
            missing.append("tesseract-ocr-eng")
        return {"ready": not missing, "missing": missing, "executable": executable, "languages": sorted(languages)}

    async def extract_text(self, file_path: str, file_type: str) -> str:
        """提取文档文本"""
        file_type_lower = file_type.lower()

        if file_type_lower in ["pdf", ".pdf"]:
            return self._extract_pdf(file_path)
        elif file_type_lower in ["txt", ".txt", "text"]:
            return self._extract_txt(file_path)
        elif file_type_lower in ["md", ".md", "markdown"]:
            return self._extract_txt(file_path)
        elif file_type_lower in ["docx", ".docx"]:
            return self._extract_docx(file_path)
        elif file_type_lower in ["html", ".html", "htm"]:
            return self._extract_html(file_path)
        elif file_type_lower in ["epub", ".epub"]:
            return self._extract_epub(file_path)
        elif file_type_lower in ["xlsx", ".xlsx", "xls", ".xls"]:
            return self._extract_xlsx(file_path)
        elif file_type_lower in ["csv", ".csv"]:
            return self._extract_csv(file_path)
        elif file_type_lower in ["image", ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"]:
            return self._extract_image(file_path)
        else:
            return self._extract_txt(file_path)

    def _extract_pdf(self, file_path: str) -> str:
        """提取PDF文本"""
        try:
            # 运行镜像目前固定安装 pypdf2（导入名 PyPDF2），部分开发环境
            # 则提供新包 pypdf；两者 API 兼容，不能因为包名差异静默返回空文本。
            try:
                from pypdf import PdfReader
                parser_name = "pypdf"
            except ImportError:
                from PyPDF2 import PdfReader
                parser_name = "PyPDF2"
            reader = PdfReader(file_path)
            text_parts = []
            for page_index, page in enumerate(reader.pages, 1):
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    text_parts.append(f"【第 {page_index} 页】\n{page_text}")
            text = "\n\n".join(text_parts).strip()
            logger.info("PDF 文本提取完成 parser=%s pages=%s chars=%s", parser_name, len(reader.pages), len(text))
            return text
        except (ImportError, OSError, RuntimeError, ValueError, TypeError) as exc:
            logger.warning("PDF 文本提取失败 file=%s error=%s", file_path, exc)
            return ""

    def _extract_txt(self, file_path: str) -> str:
        """提取文本文件"""
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
        except (OSError, UnicodeDecodeError, RuntimeError, ValueError, TypeError):
            try:
                with open(file_path, "r", encoding="gbk") as f:
                    return f.read()
            except (OSError, UnicodeDecodeError, RuntimeError, ValueError, TypeError):
                return ""

    def _extract_image(self, file_path: str) -> str:
        """通过 OCR 提取截图/图片中的文字；缺少 OCR 运行时则明确记录失败。"""
        try:
            from PIL import Image
            import pytesseract

            image = Image.open(file_path)
            try:
                text = pytesseract.image_to_string(image, lang="chi_sim+eng")
            except pytesseract.TesseractError:
                text = pytesseract.image_to_string(image, lang="eng")
            text = text.strip()
            logger.info("图片 OCR 完成 file=%s size=%s chars=%s", file_path, image.size, len(text))
            return text
        except (ImportError, OSError, RuntimeError, ValueError, TypeError) as exc:
            logger.warning("图片 OCR 失败 file=%s error=%s", file_path, exc)
            return ""

    def _extract_docx(self, file_path: str) -> str:
        """提取Word文档"""
        try:
            from docx import Document
            doc = Document(file_path)
            return "\n\n".join([para.text for para in doc.paragraphs])
        except (ImportError, OSError, RuntimeError, ValueError, TypeError):
            return ""

    def _extract_html(self, file_path: str) -> str:
        """提取HTML文本"""
        try:
            from bs4 import BeautifulSoup
            with open(file_path, "r", encoding="utf-8") as f:
                soup = BeautifulSoup(f.read(), "html.parser")
            return soup.get_text(separator="\n\n", strip=True)
        except (ImportError, OSError, RuntimeError, ValueError, TypeError):
            return ""

    def _extract_epub(self, file_path: str) -> str:
        """提取EPUB文本"""
        try:
            return self._extract_epub_with_ebooklib(file_path)
        except (ImportError, OSError, RuntimeError, ValueError, TypeError, AttributeError) as e:
            logger.warning(f"ebooklib 解析失败，尝试备用方案: {e}")
            try:
                return self._extract_epub_with_zipfile(file_path)
            except (OSError, RuntimeError, ValueError, TypeError, AttributeError) as e2:
                logger.error(f"备用方案也失败: {e2}", exc_info=True)
                return ""
    
    def _extract_epub_with_ebooklib(self, file_path: str) -> str:
        """使用 ebooklib 提取 EPUB 文本"""
        import ebooklib
        from ebooklib import epub
        from ebooklib.utils import debug
        from bs4 import BeautifulSoup
        debug(False)
        
        logger.info(f"使用 ebooklib 解析 EPUB 文件: {file_path}")
        
        book = epub.read_epub(file_path, options={'ignore_ncx': True})
        text_parts = []
        
        # 先尝试获取书名
        if book.title:
            text_parts.append(book.title)
        
        # 遍历所有文档项目
        for item in book.get_items():
            try:
                if item.get_type() == ebooklib.ITEM_DOCUMENT:
                    content = item.content
                    if content:
                        soup = BeautifulSoup(content, "html.parser")
                        # 移除 script 和 style 标签
                        for script in soup(["script", "style"]):
                            script.decompose()
                        text = soup.get_text(separator="\n", strip=True)
                        if text.strip():
                            text_parts.append(text)
            except (RuntimeError, ValueError, TypeError, AttributeError) as e:
                logger.warning(f"处理 EPUB 项目时出错: {e}", exc_info=True)
                continue
        
        full_text = "\n\n".join(text_parts)
        logger.info(f"ebooklib 解析完成，提取了 {len(full_text)} 个字符")
        
        if not full_text or not full_text.strip():
            raise ValueError("未提取到文本内容")
        
        return full_text
    
    def _extract_epub_with_zipfile(self, file_path: str) -> str:
        """使用 zipfile 直接解压并解析 EPUB（备用方案）"""
        import zipfile
        from bs4 import BeautifulSoup
        logger.info(f"使用 zipfile 解析 EPUB 文件: {file_path}")
        
        text_parts = []
        
        with zipfile.ZipFile(file_path, 'r') as zf:
            # 获取所有文件
            all_files = zf.namelist()
            
            # 优先查找 HTML 和 XHTML 文件
            html_files = [f for f in all_files if f.lower().endswith(('.html', '.xhtml', '.htm'))]
            
            # 按文件名排序，保持内容顺序
            html_files.sort()
            
            for html_file in html_files:
                try:
                    with zf.open(html_file) as f:
                        content = f.read()
                        soup = BeautifulSoup(content, "html.parser")
                        # 移除 script 和 style 标签
                        for script in soup(["script", "style"]):
                            script.decompose()
                        text = soup.get_text(separator="\n", strip=True)
                        if text.strip():
                            text_parts.append(text)
                except (OSError, RuntimeError, ValueError, TypeError, AttributeError) as e:
                    logger.warning(f"处理文件 {html_file} 时出错: {e}", exc_info=True)
                    continue
        
        full_text = "\n\n".join(text_parts)
        logger.info(f"zipfile 解析完成，提取了 {len(full_text)} 个字符")
        
        return full_text

    def _extract_xlsx(self, file_path: str) -> str:
        """提取Excel表格文本（支持 .xlsx 和旧版 .xls）"""
        # 优先尝试 openpyxl（.xlsx）
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file_path, read_only=True, data_only=True)
            all_parts = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                sheet_parts = [f"【工作表：{sheet_name}】"]
                for row in ws.iter_rows(values_only=True):
                    row_values = [str(cell) if cell is not None else "" for cell in row]
                    row_text = "\t".join(row_values).strip()
                    if row_text:
                        sheet_parts.append(row_text)
                if len(sheet_parts) > 1:
                    all_parts.append("\n".join(sheet_parts))
            wb.close()
            return "\n\n".join(all_parts)
        except (ImportError, OSError, RuntimeError, ValueError, TypeError):
            pass

        # 旧版 .xls 回退到 xlrd
        try:
            import xlrd
            wb = xlrd.open_workbook(file_path)
            all_parts = []
            for sheet_name in wb.sheet_names():
                ws = wb.sheet_by_name(sheet_name)
                sheet_parts = [f"【工作表：{sheet_name}】"]
                for row_idx in range(ws.nrows):
                    row_values = []
                    for col_idx in range(ws.ncols):
                        cell_val = ws.cell_value(row_idx, col_idx)
                        row_values.append(str(cell_val) if cell_val != "" else "")
                    row_text = "\t".join(row_values).strip()
                    if row_text:
                        sheet_parts.append(row_text)
                if len(sheet_parts) > 1:
                    all_parts.append("\n".join(sheet_parts))
            return "\n\n".join(all_parts)
        except (ImportError, OSError, RuntimeError, ValueError, TypeError):
            return ""

    def _extract_csv(self, file_path: str) -> str:
        """提取CSV文本"""
        try:
            import csv
            with open(file_path, "r", encoding="utf-8") as f:
                reader = csv.reader(f)
                rows = ["\t".join(row) for row in reader if any(cell.strip() for cell in row)]
            return "\n".join(rows)
        except (OSError, UnicodeDecodeError, RuntimeError, ValueError, TypeError):
            try:
                with open(file_path, "r", encoding="gbk") as f:
                    reader = csv.reader(f)
                    rows = ["\t".join(row) for row in reader if any(cell.strip() for cell in row)]
                return "\n".join(rows)
            except (OSError, UnicodeDecodeError, RuntimeError, ValueError, TypeError):
                return ""

    def supports(self, file_type: str) -> bool:
        """检查是否支持该文件类型"""
        supported = ["pdf", "txt", "md", "docx", "html", "epub", "image",
                     "png", "jpg", "jpeg", "webp", "bmp", "gif",
                     "xlsx", "xls", "csv"]
        return file_type.lower().replace(".", "") in supported

    def split_text(self, text: str, chunk_size: Optional[int] = None, chunk_overlap: Optional[int] = None) -> List[str]:
        """分块文本"""
        size = chunk_size or self.chunk_size
        overlap = chunk_overlap or self.chunk_overlap
        return self.text_splitter.split(text, size, overlap)

    def create_chunks(
        self,
        text: str,
        document_id: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> List[Dict[str, Any]]:
        """创建文档块"""
        chunks = self.split_text(text)
        result = []

        for i, chunk in enumerate(chunks):
            result.append({
                "id": f"{document_id}_chunk_{i}",
                "document_id": document_id,
                "chunk_index": i,
                "content": chunk,
                "metadata": metadata or {}
            })

        return result
